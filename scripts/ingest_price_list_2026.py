"""
Parse the 2026-07-01 Restar Solar Cameroon price list, apply the 20% margin
markup (protects the local distributor's margin — see repo history for the
"AI quoted prices too low" incident that prompted this), and embed bilingual
(EN/FR) knowledge-base chunks into the existing ChromaDB collection.

Only the marked-up figures are ever embedded — the pre-markup base prices
from the source spreadsheet are read in memory to compute the markup and are
never written into any chunk text, so the AI can only ever quote the
distributor-safe number.

Source: data/price_list/restar_prices_cameroon_20260701_base.xlsx
Idempotent (upsert by deterministic doc_id).

Run: python scripts/ingest_price_list_2026.py
"""
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

from openpyxl import load_workbook

from app.rag.embedder import embed_batch
from app.rag.retriever import upsert

SOURCE_XLSX = Path("data/price_list/restar_prices_cameroon_20260701_base.xlsx")
MARKUP = 1.2
EFFECTIVE_DATE = "2026-07-01"
SOURCE_FILE_TAG = "restar_price_list_20260701"

CATEGORY_LABEL = {
    "Solar panel": ("Solar Panel", "Panneau Solaire"),
    "Lithium Battery": ("Lithium Battery", "Batterie au Lithium"),
    "GEL battery": ("GEL Battery", "Batterie GEL"),
    "AGM battery": ("AGM Battery", "Batterie AGM"),
    "off-gird Inverter": ("Off-Grid Inverter", "Onduleur Hors Réseau"),
    "Inverter": ("Inverter", "Onduleur"),
    "AT controller": ("AT Charge Controller", "Régulateur de Charge AT"),
    "NT controller": ("NT Charge Controller", "Régulateur de Charge NT"),
    "CT controller": ("CT Charge Controller", "Régulateur de Charge CT"),
    "MPPT controller": ("MPPT Charge Controller", "Régulateur de Charge MPPT"),
    "Frzeer": ("Freezer", "Congélateur"),
    "Refrigerator": ("Refrigerator", "Réfrigérateur"),
    "Road Light": ("Solar Road Light", "Lampadaire Solaire"),
    "Pole for ST Light": ("Street Light Pole", "Poteau pour Lampadaire"),
    "Ceiling fan": ("Ceiling Fan", "Ventilateur de Plafond"),
    "Solar fan": ("Solar Fan", "Ventilateur Solaire"),
    "Cable": ("Cable", "Câble"),
}


def _category_label(raw: str):
    key = raw.strip()
    if key.startswith("Flood Light"):
        return ("Solar Flood Light", "Projecteur Solaire")
    if key in CATEGORY_LABEL:
        return CATEGORY_LABEL[key]
    return (key, key)


def marked_up(value):
    """Apply MARKUP to a numeric cell, or to both numbers in an 'A--B' / 'A-B'
    range string. Returns None if the cell isn't a real price."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return round(value * MARKUP)
    text = str(value).strip()
    m = re.match(r"^(\d+)\s*-{1,2}\s*(\d+)$", text)
    if m:
        lo, hi = int(m.group(1)), int(m.group(2))
        return f"{round(lo * MARKUP)}-{round(hi * MARKUP)}"
    return None


DISPLAY_FIX = {
    "Frzeer": "Freezer",
    "off-gird Inverter": "Off-Grid Inverter",
}


def parse_rows():
    """Yields (top_category_en, top_category_fr, line_label, spec, [(tier_label, marked_up_price), ...]).

    top_category is the coarse section (from the last 'Model' header row) —
    used only for metadata grouping. line_label is the row's own first
    column, which is the actual distinguishing name for that line (a specific
    panel model for Solar panel rows, or the repeated section name for
    sections like Lithium Battery/Cable where column 2 carries the real
    spec instead)."""
    wb = load_workbook(SOURCE_XLSX, data_only=True)
    ws = wb.active
    tier_labels = []
    top_category = ("", "")
    for row in ws.iter_rows(min_row=1, values_only=True):
        col1, col2 = row[1], row[2]
        if col1 is None:
            continue
        if col2 == "Model":
            # Header row: col1 is the section name, columns 3.. are this
            # section's tier labels.
            top_category = _category_label(col1)
            tier_labels = [str(v).strip() for v in row[3:9] if v is not None]
            continue
        if not tier_labels:
            continue  # rows before the first header (title/date) — skip

        raw_prices = row[3:6]
        tiers = []
        for label, raw in zip(tier_labels, raw_prices):
            adj = marked_up(raw)
            if adj is not None:
                tiers.append((label, adj))
        if not tiers:
            continue  # e.g. the "Pole for ST Light" freeform note row

        line_label = DISPLAY_FIX.get(str(col1).strip(), str(col1).strip())
        spec = str(col2).strip() if col2 is not None else ""
        yield top_category[0], top_category[1], line_label, spec, tiers


def chunk_en(line_label: str, spec: str, tiers: list) -> str:
    price_bits = "; ".join(f"{label}: {price} FCFA each" for label, price in tiers)
    spec_bit = f" ({spec})" if spec and spec != "Model" else ""
    return (
        f"Q: How much does the Restar Solar {line_label}{spec_bit} cost? "
        f"A: The {line_label}{spec_bit} — {price_bits}. "
        f"Prices effective {EFFECTIVE_DATE}, Douala, Cameroon, subject to change without notice."
    )


def chunk_fr(line_label: str, spec: str, tiers: list) -> str:
    price_bits = "; ".join(f"{label} : {price} FCFA/pièce" for label, price in tiers)
    spec_bit = f" ({spec})" if spec and spec != "Model" else ""
    return (
        f"Q: Combien coûte le Restar Solar {line_label}{spec_bit} ? "
        f"R : Le {line_label}{spec_bit} — {price_bits}. "
        f"Prix en vigueur au {EFFECTIVE_DATE}, Douala, Cameroun, sous réserve de modification."
    )


def main():
    if not SOURCE_XLSX.exists():
        raise SystemExit(f"Missing source file: {SOURCE_XLSX}")

    doc_ids, texts, metadatas = [], [], []

    for idx, (cat_en, cat_fr, line_label, spec, tiers) in enumerate(parse_rows()):
        base_id = f"price_{SOURCE_FILE_TAG}_{idx}"

        doc_ids.append(f"{base_id}_en")
        texts.append(chunk_en(line_label, spec, tiers))
        metadatas.append({
            "language": "en", "category": cat_en, "source_file": SOURCE_FILE_TAG,
            "kind": "price_list", "effective_date": EFFECTIVE_DATE,
        })

        doc_ids.append(f"{base_id}_fr")
        texts.append(chunk_fr(line_label, spec, tiers))
        metadatas.append({
            "language": "fr", "category": cat_fr, "source_file": SOURCE_FILE_TAG,
            "kind": "price_list", "effective_date": EFFECTIVE_DATE,
        })

    print(f"Parsed {len(texts) // 2} priced line items from {SOURCE_XLSX.name}.")
    print(f"Embedding {len(texts)} chunks (local sentence-transformers model, may take a minute)...")
    embeddings = embed_batch(texts)
    for doc_id, text, vec, meta in zip(doc_ids, texts, embeddings, metadatas):
        upsert(doc_id, text, vec, meta)

    print(f"Upserted {len(texts)} price-list knowledge-base chunks into ChromaDB.")


if __name__ == "__main__":
    main()
