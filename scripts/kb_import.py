"""
Import a document into the Rest Solar knowledge base.

Usage:
    python scripts/kb_import.py path/to/file.pdf
    python scripts/kb_import.py path/to/prices.xlsx
    python scripts/kb_import.py path/to/page.html
    python scripts/kb_import.py path/to/notes.txt

Supported: .pdf  .xlsx  .html  .htm  .txt
"""
import sys
import re
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

from dotenv import load_dotenv
load_dotenv()

from app.rag.embedder import embed_batch
from app.rag.retriever import upsert

CHUNK_SIZE = 500
CHUNK_OVERLAP = 80


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(str(path), read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                parts = [str(c) for c in row if c is not None]
                if parts:
                    lines.append("\t".join(parts))
        return "\n".join(lines)
    elif suffix in (".html", ".htm"):
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser")
        return soup.get_text(separator="\n")
    elif suffix == ".txt":
        return path.read_text(encoding="utf-8")
    else:
        raise ValueError(f"Unsupported file type: {suffix}")


def chunk_text(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    chunks, start = [], 0
    while start < len(text):
        end = min(start + chunk_size, len(text))
        boundary = text.rfind(". ", start, end)
        if boundary != -1 and boundary > start + chunk_size // 2:
            end = boundary + 1
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end == len(text):
            break
        start = end - overlap
    return chunks


def main(file_path: str):
    path = Path(file_path).resolve()
    if not path.exists():
        print(f"Error: file not found: {path}", file=sys.stderr)
        sys.exit(1)

    print(f"Importing: {path.name}")
    text = extract_text(path)
    if not text.strip():
        print("Warning: no text extracted from file.")
        return

    chunks = chunk_text(text)
    print(f"  -> {len(chunks)} chunks")

    embeddings = embed_batch(chunks)
    for i, (chunk, vec) in enumerate(zip(chunks, embeddings)):
        doc_id = f"{path.stem}_{i}"
        upsert(doc_id, chunk, vec, {
            "source_file": path.name,
            "language": "unknown",
            "chunk_index": i,
        })

    print(f"Imported {len(chunks)} chunks from {path.name} into ChromaDB")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python scripts/kb_import.py <file>", file=sys.stderr)
        sys.exit(1)
    main(sys.argv[1])
